"""
Turns the Acquisition Checklist workbook into the two files the app needs:

  src/data/brands/jm-acquisitions.json     the checklist itself
  scripts/seed-jm-acquisitions-projects.json   the stores, for the importer

The workbook is shaped differently from the brand workbooks: Jersey Mike's
and Dave's are one column per checklist item, where this is one ROW per item
with its own priority, owner and notes. So each row becomes one checklist
item, and the per-row columns become the item's metadata:

  Item                  -> the label
  In-Charge             -> resp, shown after the label like an assignee
  Priority              -> a dropdown, its options taken from the sheet's own
                           data validation rather than from the values used,
                           so an option nobody has picked yet still appears
  Done                  -> the tick-box itself, and what progress counts
  Vendor / Notes /
  How To / SLA /
  Contact / Progress    -> a note under the item

Usage:  python3 scripts/convert-acquisition-workbook.py <workbook.xlsx>
"""
import json
import re
import sys
from openpyxl import load_workbook

SRC = sys.argv[1] if len(sys.argv) > 1 else sys.exit('pass the .xlsx path')

# The note columns differ between the two sheets; both are "things worth
# knowing while doing this item", so they are joined into one note.
NOTE_COLS = {
    'Pre Close': ['Vendor', 'Notes', 'How To'],
    'Post Close': ['Vendor', 'SLA', 'Contact', 'Progress Notes'],
}
PREFIX = {'Pre Close': 'PRE', 'Post Close': 'POST'}


def clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def options_for(ws, col_letter):
    """The dropdown list Excel itself enforces on a column."""
    for dv in ws.data_validations.dataValidation:
        if dv.type != 'list':
            continue
        if any(str(r).startswith(f'{col_letter}') for r in dv.sqref.ranges):
            raw = (dv.formula1 or '').strip('"')
            seen, out = set(), []
            for o in (x.strip() for x in raw.split(',')):
                if o and o not in seen:
                    seen.add(o)
                    out.append(o)
            return out
    return []


wb = load_workbook(SRC)
wbv = load_workbook(SRC, data_only=True)

headers = []
# Reference data about the store, kept at the top of the Pre Close list.
# 'C' deliberately matches the Jersey Mike's address column, because the
# project list and the search box read the address from C.
headers.append({'col': 2, 'letter': 'B', 'label': 'Phone Number', 'phase': 'Pre Close',
                'resp': None, 'hint': None, 'type': 'text'})
headers.append({'col': 3, 'letter': 'C', 'label': 'Full Address', 'phase': 'Pre Close',
                'resp': None, 'hint': None, 'type': 'text'})

col = 4
report = {}
for sheet in ['Pre Close', 'Post Close']:
    ws, wsv = wb[sheet], wbv[sheet]
    hdr = [clean(c) for c in next(wsv.iter_rows(min_row=2, max_row=2, values_only=True))]
    idx = {h: i for i, h in enumerate(hdr) if h}
    opts = options_for(ws, 'B')
    n = 0
    dues = 0
    for row in wsv.iter_rows(min_row=3, values_only=True):
        get = lambda k: clean(row[idx[k]]) if k in idx and idx[k] < len(row) else ''
        item = get('Item')
        if not item:
            continue  # filler rows carrying only Done=False
        n += 1
        if get('Due Date'):
            dues += 1
        note = ' · '.join(x for x in (get(c) for c in NOTE_COLS[sheet]) if x)
        priority = get('Priority')
        headers.append({
            'col': col,
            'letter': f'{PREFIX[sheet]}{n}',
            'label': item,
            'phase': sheet,
            'resp': get('In-Charge') or None,
            'hint': note or None,
            'type': 'checkbox',
            'options': opts,
            # The sheet's own priority for this item, used until someone
            # changes it on a store. Not every row has one.
            'priority': priority if priority in opts else None,
        })
        col += 1
    report[sheet] = {'items': n, 'options': opts, 'rows_with_due_date': dues}

with open('src/data/brands/jm-acquisitions.json', 'w') as f:
    json.dump(headers, f, indent=1)
    f.write('\n')

# --- the stores -------------------------------------------------------
ws = wbv['Stores']
stores = []
for i, row in enumerate(ws.iter_rows(min_row=4, values_only=True)):
    name, phone, addr = clean(row[2]), clean(row[3]), clean(row[4])
    if not name:
        continue
    # Most stores are "6030- Mt Vernon, VA". The few with no store number
    # yet are left as "-Seminole Trail", which reads as a typo and sorts
    # ahead of everything. The separator is only meaningful after a number.
    name = re.sub(r'^[\s\-]+', '', name)
    fields = {}
    if phone:
        fields['B'] = phone
    if addr:
        fields['C'] = addr
    stores.append({
        'row': i + 4,
        'brand': "Jersey Mike's",
        'brandKey': 'jmacquisitions',
        'name': name,
        'order': len(stores),
        'completed': False,
        'fields': fields,
    })

with open('scripts/seed-jm-acquisitions-projects.json', 'w') as f:
    json.dump(stores, f, indent=1)
    f.write('\n')

print(json.dumps(report, indent=2))
print(f'\n{len(headers)} checklist items written (2 reference fields + tasks)')
print(f'{len(stores)} stores written')
print('stores missing an address:',
      [s['name'] for s in stores if 'C' not in s['fields']] or 'none')
